"""Offline end-to-end fixtures: tiny PDFs, local OCR JSON and PNG assets only."""
from __future__ import annotations

import copy
import io
import json
import re
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from PIL import Image
from markdown_it import MarkdownIt

ROOT = (Path(__file__).resolve().parents[2] / 'skills' / 'sanmu-doc2md')
sys.path.insert(0, str(ROOT / 'scripts'))
import build_markdown as build
from build_support import OWNERSHIP, Publisher, asset_path, load_book


def block(kind='text', text='', page=0, bbox=None, **fields):
    return {'type': kind, 'text': text, 'page_idx': page,
            'bbox': bbox or [100, 200, 600, 400], **fields}


def png(color):
    buffer = io.BytesIO()
    Image.new('RGB', (12, 10), color).save(buffer, format='PNG')
    return buffer.getvalue()


class BuildTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='test-build-')
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.raw = self.base / 'raw'
        self.raw.mkdir()
        self.out = self.base / 'out'
        self.manifest_path = self.base / 'manifest.json'
        self.chapters_path = self.base / 'chapters.json'

    def write_json(self, path, data):
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')

    def part(self, identity, items, *, nested=False, color='red'):
        folder = self.raw / identity
        ocr = folder / 'result' if nested else folder
        self.write_json(ocr / 'book_content_list.json', items)
        (ocr / 'images' / 'nested').mkdir(parents=True, exist_ok=True)
        (ocr / 'images' / 'same.png').write_bytes(png(color))
        (ocr / 'images' / 'nested' / 'qr.png').write_bytes(png(color))
        return ocr

    def book(self, *, total=2, pdf=False):
        manifest = {'pages': total, 'parts': [
            {'id': f'part{i:02}', 'start_page': i, 'end_page': i, 'pages': 1}
            for i in range(1, total + 1)]}
        sections = [{'id': f'ch{i:02}', 'title': f'传入标题{i}', 'start': i, 'end': i}
                    for i in range(1, total + 1)]
        if pdf:
            path = self.base / 'source.pdf'
            with fitz.open() as doc:
                for i in range(total):
                    page = doc.new_page(width=200, height=100)
                    color = [(1, 0, 0), (0, 0, 1)][i % 2]
                    page.draw_rect(page.rect, color=color, fill=color)
                doc.save(path)
            manifest['source'] = str(path)
        self.write_json(self.manifest_path, manifest)
        self.write_json(self.chapters_path, sections)
        return manifest, sections

    def cli(self, *args, book=False, ok=True):
        command = [sys.executable, str(ROOT / 'scripts' / 'build_markdown.py'),
                   '--raw', str(self.raw), '--out', str(self.out)]
        if book:
            command += ['--manifest', str(self.manifest_path), '--chapters', str(self.chapters_path)]
        result = subprocess.run(command + list(args), capture_output=True, text=True,
                                encoding='utf-8', errors='replace', timeout=30)
        if ok:
            self.assertEqual(result.returncode, 0, result.stderr + result.stdout)
        else:
            self.assertNotEqual(result.returncode, 0, result.stderr + result.stdout)
        return result

    def text(self, name='ch01.md'):
        root = self.out.parent / f'.{self.out.name}.build' if name in ('build-report.json', OWNERSHIP) else self.out
        return (root / name).read_text(encoding='utf-8')

    def refs(self, text):
        return re.findall(r'!\[[^\]]*\]\(([^)]+)\)', text)

    def render(self, items, **kwargs):
        title = kwargs.pop('title', '传入标题')
        return build.build_chapter(items, None, kwargs.pop('pdf_path', None),
                                   self.base / 'render' / 'images' / 'ch', title, [],
                                   mode='book', **kwargs)

    def snapshot(self):
        return {p.relative_to(self.out).as_posix(): p.read_bytes()
                for p in self.out.rglob('*') if p.is_file()} if self.out.exists() else {}

    def test_pdf_page_offset_and_crop_coordinates(self):
        self.book(pdf=True)
        for i in (1, 2):
            self.part(f'part{i:02}', [block('image', img_path='images/same.png',
                                          bbox=[100, 200, 500, 700])])
        self.cli(book=True)
        for i, expected in [(1, (255, 0, 0)), (2, (0, 0, 255))]:
            ref = self.refs(self.text(f'ch{i:02}.md'))[0]
            with Image.open(self.out / ref) as image:
                self.assertAlmostEqual(image.width / image.height, 1.6, places=2)
                pixel = image.convert('RGB').getpixel((image.width // 2, image.height // 2))
                self.assertTrue(all(abs(a - b) < 3 for a, b in zip(pixel, expected)), pixel)

    def test_part_same_image_names_nested_tables_and_order(self):
        self.book()
        for i, color in [(1, 'red'), (2, 'blue')]:
            self.part(f'part{i:02}', [
                block(text=f'前正文{i}'),
                block('image', img_path='images/same.png', image_caption=[f'图{i} 图注']),
                block(text=f'后正文{i}'),
                block('table', table_caption=f'表{i} 表题',
                      table_body='<table><tr><td>单元格<img src="images/nested/qr.png" alt="二维码"></td></tr></table>',
                      table_footnote='整条表下注'),
                block('page_footnote', text=f'脚注{i}')
            ], nested=True, color=color)
        self.cli(book=True)
        for i, color in [(1, (255, 0, 0)), (2, (0, 0, 255))]:
            text = self.text(f'ch{i:02}.md')
            self.assertLess(text.index(f'前正文{i}'), text.index('!['))
            self.assertLess(text.index('!['), text.index(f'后正文{i}'))
            self.assertIn(f'表{i} 表题', text)
            self.assertIn(f'脚注{i}', text)
            self.assertIn('整条表下注', text)
            self.assertIn('![二维码]', text)
            for ref in self.refs(text):
                with Image.open(self.out / ref) as image:
                    self.assertEqual(image.getpixel((0, 0)), color)
        self.assertNotEqual(self.refs(self.text())[0], self.refs(self.text('ch02.md'))[0])

    def test_duplicate_references_keep_file_and_each_caption(self):
        ocr = self.part('part01', [])
        items = [block('image', img_path='images/same.png', image_caption=['图1 首次']),
                 block(text='中间正文'),
                 block('image', img_path='images/same.png', image_caption=['图2 再次'])]
        text, stats = self.render(items, ocr_dir=ocr)
        refs = self.refs(text)
        self.assertEqual(len(refs), 2)
        self.assertEqual(refs[0], refs[1])
        self.assertEqual(len(list((self.base / 'render').rglob('*.png'))), 1)
        self.assertIn('图1 首次', text)
        self.assertIn('图2 再次', text)
        self.assertEqual(stats['figures'], 2)

    def test_union_crop_does_not_mutate_items_or_skip_following_text(self):
        _, _ = self.book(total=1, pdf=True)
        items = [block('image', bbox=[100, 200, 500, 700], img_path='a.png'),
                 block('image', bbox=[300, 200, 800, 700], img_path='b.png'),
                 block(text='紧随图后的正文')]
        before = copy.deepcopy(items)
        text, _ = self.render(items, pdf_path=self.base / 'source.pdf')
        self.assertEqual(items, before)
        self.assertIn('紧随图后的正文', text)
        self.assertEqual(len(self.refs(text)), 1)
        with Image.open(self.base / 'render' / self.refs(text)[0]) as image:
            self.assertAlmostEqual(image.width / image.height, 2.8, places=2)

    def test_image_groups_respect_explicit_boundaries_on_either_side(self):
        for marked in (0, 1, 2):
            with self.subTest(marked=marked):
                items = [block('image', img_path=f'{i}.png',
                               bbox=[100 + i * 50, 200, 600 + i * 50, 700])
                         for i in range(3)]
                items[marked]['_figure_boundary'] = True
                groups = list(build.image_groups(items).values())
                self.assertIn([items[marked]], groups)
                self.assertEqual([item for group in groups for item in group], items)
                self.assertEqual([len(group) for group in groups],
                                 [1, 1, 1] if marked == 1 else ([1, 2] if marked == 0 else [2, 1]))

    def test_image_groups_block_conflicting_figure_numbers(self):
        for left, right in [('图1-2 左图', '图 1－3 右图'),
                            ('Figure 1-2. Left', 'Fig. 1-3 Right'),
                            ('FIGURE 12 Left', 'fig. 13 Right'),
                            ('图1-2 左图', 'Figure 1-3 Right')]:
            with self.subTest(left=left, right=right):
                items = [block('image', img_path='a.png', bbox=[100, 200, 600, 700],
                               image_caption=[left]),
                         block('chart', img_path='b.png', bbox=[150, 200, 650, 700]),
                         block('image', img_path='c.png', bbox=[200, 200, 700, 700],
                               image_caption=[right])]
                self.assertEqual(list(build.image_groups(items).values()), [items[:2], items[2:]])

    def test_image_groups_keep_existing_overlap_rules_and_untitled_behavior(self):
        for left, right in [('图1-2 左图', '图 1－2 右图'),
                            ('Figure 1-2 Left', 'Fig. 1-2 Right'),
                            ('图1-2 左图', 'Figure 1-2 Right'),
                            ('', '图1-2 图题'), ('图1-2 图题', ''), ('', '')]:
            with self.subTest(left=left, right=right):
                items = [block('image', img_path='a.png', bbox=[100, 200, 600, 700],
                               image_caption=[left]),
                         block('chart', img_path='b.png', bbox=[150, 200, 650, 700],
                               chart_caption=right)]
                self.assertEqual(list(build.image_groups(items).values()), [items])
                items[1]['bbox'] = [700, 200, 900, 700]
                self.assertEqual(list(build.image_groups(items).values()), [[items[0]], [items[1]]])
                items[1]['bbox'] = items[0]['bbox']
                self.assertEqual(list(build.image_groups(items).values()), [[items[0]], [items[1]]])
                items[1]['bbox'] = [150, 200, 650, 700]
                items[1]['img_path'] = 'a.png'
                self.assertEqual(list(build.image_groups(items).values()), [[items[0]], [items[1]]])

    def test_crop_failure_falls_back_to_ocr(self):
        self.book(total=1, pdf=True)
        ocr = self.part('part01', [])
        with patch.object(build.FigureStore, 'crop', return_value=None):
            text, stats = self.render([block('image', img_path='images/same.png')],
                                      ocr_dir=ocr, pdf_path=self.base / 'source.pdf')
        self.assertEqual(len(self.refs(text)), 1)
        self.assertEqual(stats['ocr_figures'], 1)

    def test_repaired_images_require_successful_pdf_crop(self):
        self.book(total=1, pdf=True)
        ocr = self.part('part01', [])
        for marker in ('_recrop_required', '_added_image'):
            with self.subTest(marker=marker):
                item = block('image', img_path='images/same.png', **{marker: True})
                with self.assertRaisesRegex(ValueError, '必须从原 PDF 重裁'):
                    self.render([item], ocr_dir=ocr)
                with patch.object(build.FigureStore, 'crop', return_value=None):
                    with self.assertRaisesRegex(ValueError, 'PDF 重裁失败'):
                        self.render([item], ocr_dir=ocr, pdf_path=self.base / 'source.pdf')

    def test_repaired_expanded_bbox_crops_new_pdf_pixels_instead_of_reusing_native_image(self):
        from book_features import apply_repairs
        pdf_path = self.base / 'edge.pdf'
        with fitz.open() as doc:
            page = doc.new_page(width=200, height=100)
            page.draw_rect(page.rect, color=(1, 0, 0), fill=(1, 0, 0))
            page.draw_rect(fitz.Rect(140, 20, 180, 80), color=(0, 0, 1), fill=(0, 0, 1))
            doc.save(pdf_path)
        ocr = self.part('part01', [], color='green')
        item = block('image', bbox=[100, 200, 500, 800], img_path='images/same.png',
                     image_caption=['图1 原图'], _page=1)
        old = {field: item.get(field, []) for field in
               ('type', 'bbox', 'img_path', 'image_caption', 'image_footnote')}
        repaired, _ = apply_repairs([item], [{
            'page': 1, 'bbox': item['bbox'], 'action': 'update_image', 'old': old,
            'new_bbox': [100, 200, 900, 800], 'reason': '原 PDF 边缘有蓝色标记'
        }], pdf_path)
        self.assertTrue(repaired[0]['_recrop_required'])
        source = [{'texts': [], 'pictures': [{'box': (0.1, 0.2, 0.9, 0.8),
                                              'blob': png('green'), 'ext': 'png'}]}]
        image_dir = self.base / 'render' / 'images' / 'ch'
        with patch.object(build, 'best_picture_match', side_effect=AssertionError('must recrop')):
            text, stats = build.build_chapter(repaired, source, pdf_path, image_dir,
                                             '传入标题', [], ocr_dir=ocr, mode='book')
        self.assertEqual(stats['cropped_figures'], 1)
        self.assertEqual(stats.get('native_figures', 0), 0)
        self.assertEqual(stats.get('ocr_figures', 0), 0)
        with Image.open(self.base / 'render' / self.refs(text)[0]) as image:
            rgb = image.convert('RGB')
            self.assertEqual(rgb.getpixel((image.width // 4, image.height // 2)), (255, 0, 0))
            self.assertEqual(rgb.getpixel((image.width * 7 // 8, image.height // 2)), (0, 0, 255))

    def test_crop_exception_returns_failure(self):
        class BrokenPage:
            @property
            def rect(self):
                raise RuntimeError('synthetic rendering failure')
        self.assertIsNone(build.FigureStore(self.base).crop(BrokenPage(), (0, 0, 1, 1)))

    def test_full_page_scan_preserved_without_suppressing_text(self):
        ocr = self.part('part01', [])
        text, stats = self.render([block('image', bbox=[0, 0, 1000, 1000], img_path='images/same.png'),
                                   block(text='短正文')], ocr_dir=ocr)
        self.assertEqual(len(self.refs(text)), 1)
        self.assertIn('短正文', text)
        self.assertEqual(stats['full_page_scans_preserved'], 1)

    def test_title_one_h1_exercises_and_page_top_prose(self):
        items = [block(text='OCR错误书名', text_level=1),
                 block(text='1 概述', text_level=1),
                 block(text='1.1 正规小节', text_level=1),
                 block(text='1. 计算这个数', text_level=1),
                 block(text='(2) 计算另一个数', text_level=1),
                 block(text='这是页顶正文', bbox=[20, 0, 900, 50]),
                 block(text='下页正文', page=1)]
        text, _ = self.render(items)
        self.assertTrue(text.startswith('# 传入标题\n'))
        self.assertEqual(re.findall(r'^# .*', text, re.M), ['# 传入标题'])
        self.assertIn('\n## 1 概述\n', text)
        self.assertIn('\n1. 计算这个数\n', text)
        self.assertIn('\n(2) 计算另一个数\n', text)
        self.assertNotIn('## 这是页顶正文', text)

    def test_book_exercise_context_and_elective_headings(self):
        text, _ = self.render([
            block(text=r'\*1.1 选读小节', text_level=1),
            block(text='小节内的次级标题', text_level=1),
            block(text='复习题'),
            block(text=r'1.1 计算 $x_1$ 的值。', text_level=1),
            block(text='① 解释原因', text_level=1),
            block(text='（一）给出条件', text_level=1),
            block(text='这是很长的题干。它不应成为标题。', text_level=1)])
        self.assertIn('## ★ 1.1 选读小节', text)
        self.assertIn('### 小节内的次级标题', text)
        self.assertIn('## 复习题', text)
        self.assertNotIn('## 1.1 计算', text)
        self.assertNotIn('## ①', text)
        self.assertNotIn('## （一）', text)
        self.assertIn('$x_1$', text)
        self.assertNotIn('## 这是很长', text)

    def test_split_chapter_title_matches_authoritative_title_only(self):
        items = [block(text='第一章', text_level=1), block(text='气体①', text_level=1),
                 block(text='本章基本要求'), block(text='完整正文')]
        text, stats = build.build_chapter(items, None, None, self.base / 'images' / 'ch',
                                         '第一章 气体', [], mode='book')
        self.assertEqual(re.findall(r'^#{1,2} .*', text, re.M), ['# 第一章 气体', '## 本章基本要求'])
        self.assertIn('完整正文', text)
        self.assertEqual(stats['duplicate_chapter_labels'], 2)
        self.assertEqual(items[0]['text'], '第一章')

    def test_substantive_footer_kept_and_running_footer_removed(self):
        review = []
        items = [block('footer', text='① 本章补充说明，含公式 $x^2$。'),
                 block('footer', text='第一章 气体'), block(text='正文'),
                 block('footer', text='第一章 气体', page=1),
                 block('header', text='2.1 热力学概论'),
                 block('header', text=r'2.10 气体的 $\Delta U$ 和 $\Delta H$'),
                 block('header', text=r'$\Delta G$')]
        text, _ = build.build_chapter(items, None, None, self.base / 'images' / 'ch',
                                     '气体', review, mode='book')
        self.assertIn('① 本章补充说明，含公式 $x^2$。', text)
        self.assertNotIn('第一章 气体', text)
        self.assertNotIn('热力学概论', text)
        self.assertNotIn('$\\Delta G$', text)
        self.assertEqual(len(review), 1)
        self.assertEqual(review[0]['kind'], 'retained_marginal_text')

    def test_default_delivery_is_clean_and_index_requires_flag(self):
        self.book(total=1)
        self.part('part01', [block(text='正文 $x$')])
        self.cli(book=True)
        self.assertEqual({p.name for p in self.out.iterdir()}, {'00-目录.md', 'ch01.md'})
        self.assertTrue((self.out.parent / '.out.build' / 'build-report.json').is_file())
        self.cli('--agent-index', book=True)
        self.assertTrue((self.out / 'agent' / 'chunks.jsonl').is_file())
        from book_features import validate_agent_index
        self.assertEqual(validate_agent_index(self.out), [])

    def test_repairs_stale_failure_preserves_delivery_and_state(self):
        self.book(total=1)
        self.part('part01', [block(text='原文')])
        self.cli(book=True)
        before = self.snapshot()
        repair = self.base / 'repairs.json'
        self.write_json(repair, [{'page': 1, 'bbox': [100, 200, 600, 400], 'field': 'text',
                                  'old': '过期', 'new': '新文', 'reason': '核对原页'}])
        self.cli('--repairs', str(repair), book=True, ok=False)
        self.assertEqual(self.snapshot(), before)
        self.write_json(repair, [{'page': 1, 'bbox': [100, 200, 600, 400], 'field': 'text',
                                  'old': '原文', 'new': '新文', 'reason': '核对原页'}])
        self.cli('--repairs', str(repair), book=True)
        self.assertIn('新文', self.text())

    def test_empty_page_mapping_and_partial_agent_rebuild(self):
        _, sections = self.book(total=2)
        self.part('part01', [block(text='第一页')])
        self.part('part02', [])
        self.cli('--agent-index', book=True)
        from book_features import validate_agent_index
        self.assertEqual(validate_agent_index(self.out), [])
        report = json.loads(self.text('build-report.json'))
        self.assertEqual(report['page_maps']['ch02'][0]['pdf_page'], 2)
        self.assertEqual(report['page_maps']['ch02'][0]['status'], 'no_output')
        self.part('part01', [block(text='更新第一页')])
        self.cli('--agent-index', '--only', 'ch01', book=True)
        self.assertEqual(validate_agent_index(self.out), [])
        self.assertIn('ch02', (self.out / 'agent' / 'manifest.json').read_text(encoding='utf-8'))

    def test_review_regressions_keep_prose_and_heading_boundaries(self):
        text, _ = self.render([block(text='8.1.1 首个深层小节'),
                               block(text='平衡', text_level=1), block(text='另一段'),
                               block(text='平衡'), block(text='8.1 计算体系的熵变。'),
                               block(text='第一章介绍了热力学基础。'),
                               block(text='习题'), block(text='附录'), block(text='1.1 常数表')])
        self.assertIn('## 8.1.1 首个深层小节', text)
        self.assertIn('另一段\n\n平衡', text)
        self.assertNotIn('## 8.1 计算', text)
        self.assertNotIn('## 第一章介绍', text)
        self.assertIn('## 1.1 常数表', text)

    def test_short_discarded_and_repeated_footnotes_preserved(self):
        text, _ = self.render([block('discarded', text='[1] 标准态为纯物质'),
                               block('footer', text='① 质量守恒'),
                               block('footer', text='① 质量守恒', page=1)])
        self.assertIn('[1] 标准态为纯物质', text)
        self.assertEqual(text.count('① 质量守恒'), 2)

    def test_equation_prose_suffix_follows_formula(self):
        text, _ = self.render([block('equation', text=r'\text{和像差（主要是球差} \varDelta r_s = x) 。')])
        self.assertLess(text.index('和像差'), text.index('\\varDelta'))
        self.assertLess(text.index('\\varDelta'), text.index('）。'))

    def test_existing_agent_index_cannot_silently_become_stale(self):
        self.book(total=1)
        self.part('part01', [block(text='原文')])
        self.cli('--agent-index', book=True)
        before = self.snapshot()
        self.part('part01', [block(text='新文')])
        self.cli(book=True, ok=False)
        self.assertEqual(self.snapshot(), before)

    def test_partial_rebuild_retains_untouched_repair_audit(self):
        self.book()
        for i in (1, 2):
            self.part(f'part{i:02}', [block(text=f'原文{i}')])
        repair = self.base / 'repairs.json'
        self.write_json(repair, [{'page': 2, 'bbox': [100, 200, 600, 400], 'field': 'text',
                                  'old': '原文2', 'new': '已校正文2', 'reason': '核对原页'}])
        self.cli('--repairs', str(repair), book=True)
        self.cli('--only', 'ch01', book=True)
        report = json.loads(self.text('build-report.json'))
        self.assertEqual(report['repairs'][0]['page'], 2)
        self.assertIn('已校正文2', self.text('ch02.md'))

    def test_paragraphs_tables_equations_not_joined(self):
        text, stats = self.render([block(text='上页未完'), block(text='下页第一段', page=1),
                                   block(text='同页另一段', page=1),
                                   block('equation', text='x+1', page=1),
                                   block(text='公式后正文', page=1),
                                   block('table', page=1, table_body='<table><tr><td>数据</td></tr></table>'),
                                   block(text='表格后正文', page=1)])
        self.assertIn('上页未完\n\n下页第一段\n\n同页另一段', text)
        self.assertIn('$$\n\n公式后正文', text)
        self.assertIn('| 数据 |', text)
        self.assertEqual(stats.get('joined_page_paragraphs', 0), 0)

    def test_inline_math_and_code_preservation(self):
        text, _ = self.render([block(text=r'数值 $\cdotkg$ 与 \(x+1\)，代码 `$\cdotkg$`'),
                               block('code', code_body=r'$\cdotkg$')])
        self.assertIn(r'$\cdot kg$', text)
        self.assertIn(r'\(x+1\)', text)
        self.assertIn(r'`$\cdotkg$`', text)
        self.assertIn('```\n' + r'$\cdotkg$' + '\n```', text)

    def test_code_fences_parse_without_swallowing_following_prose_or_images(self):
        ocr = self.part('part01', [])
        cases = [
            ('ordinary', '  print("hello")\n\n\nprint("world")\n\n',
             '  print("hello")\n\n\nprint("world")\n\n', ''),
            ('language', '```python\nprint("hello")\n```', 'print("hello")\n', 'python'),
            ('indented', '  ```python\n  x = 1\n  ```', 'x = 1\n', 'python'),
            ('relative indentation', '   ```python\n   if True:\n       x = 1\n   ```',
             'if True:\n    x = 1\n', 'python'),
            ('tildes', '~~~python\nprint("hello")\n~~~~\n', 'print("hello")\n', 'python'),
            ('nested', '````markdown\n```python\nx = 1\n```\n````',
             '```python\nx = 1\n```\n', 'markdown'),
            ('literal fences', 'before\n```\n~~~~~\n`````\nafter',
             'before\n```\n~~~~~\n`````\nafter\n', ''),
            ('unclosed', '```python\nx = 1', '```python\nx = 1\n', ''),
            ('mismatched', '```python\nx = 1\n~~~', '```python\nx = 1\n~~~\n', ''),
            ('short closer', '````python\nx = 1\n```', '````python\nx = 1\n```\n', ''),
            ('trailing content', '```python\nx = 1\n```\nkeep me',
             '```python\nx = 1\n```\nkeep me\n', ''),
            ('early closer', '```python\nx = 1\n```\nkeep me\n```',
             '```python\nx = 1\n```\nkeep me\n```\n', ''),
            ('strip one layer', '~~~~markdown\n~~~python\nx = 1\n~~~\n~~~~',
             '~~~python\nx = 1\n~~~\n', 'markdown'),
            ('backtick info', '~~~lang `option`\nx = 1\n~~~', 'x = 1\n', 'lang `option`'),
            ('empty', '```python\n```', '', 'python'),
        ]
        for kind in ('code', 'algorithm'):
            for name, source, expected, info in cases:
                for mapped in (False, True):
                    with self.subTest(kind=kind, case=name, mapped=mapped):
                        fields = {'code_body': source} if kind == 'code' else {'text': source}
                        text, _ = self.render([
                            block(kind, **fields), block(text='代码后正文。'),
                            block('image', img_path='images/same.png', image_caption=['图1 后续图片'])
                        ], ocr_dir=ocr, page_map=[] if mapped else None)
                        tokens = MarkdownIt('commonmark').parse(text)
                        fences = [token for token in tokens if token.type == 'fence']
                        self.assertEqual(len(fences), 1)
                        self.assertEqual(fences[0].content, expected)
                        self.assertEqual(fences[0].info.strip(), info)
                        children = [child for token in tokens for child in token.children or []]
                        self.assertIn('代码后正文。', [child.content for child in children if child.type == 'text'])
                        self.assertEqual(sum(child.type == 'image' for child in children), 1)

    def test_missing_part_fails_even_with_only_and_preserves_output(self):
        self.book()
        self.part('part01', [block(text='第一章')])
        self.out.mkdir()
        (self.out / 'manual.txt').write_text('人工内容', encoding='utf-8')
        before = self.snapshot()
        result = self.cli('--only', 'ch01', book=True, ok=False)
        self.assertIn('part02', result.stderr)
        self.assertEqual(self.snapshot(), before)

    def test_ranges_reject_overlap_omission_and_out_of_bounds(self):
        original, sections = self.book()
        for i in (1, 2):
            self.part(f'part{i:02}', [block(text='正文')])
        variants = []
        m = copy.deepcopy(original); m['parts'][1]['start_page'] = 1; variants.append((m, sections))
        m = copy.deepcopy(original); m['pages'] = 3; variants.append((m, sections))
        s = copy.deepcopy(sections); s[1]['end'] = 3; variants.append((original, s))
        s = copy.deepcopy(sections); s[1]['start'] = 1; variants.append((original, s))
        variants.append((original, sections[:1]))
        for manifest, chapter in variants:
            with self.subTest(manifest=manifest, sections=chapter):
                self.write_json(self.manifest_path, manifest)
                self.write_json(self.chapters_path, chapter)
                self.cli(book=True, ok=False)
                self.assertFalse(self.out.exists())

    def test_ocr_and_pdf_page_bounds_fail(self):
        self.book(total=1, pdf=True)
        self.part('part01', [block(text='越界', page=1)])
        self.cli(book=True, ok=False)
        self.assertFalse(self.out.exists())
        self.part('part01', [block(text='越界', page=-1)])
        self.cli(book=True, ok=False)
        self.part('part01', [block(text='越界', page=True)])
        self.cli(book=True, ok=False)

    def test_path_escape_inputs_and_outputs_fail(self):
        manifest, sections = self.book(total=1)
        self.part('part01', [block(text='正文')])
        for unsafe in ('../escape.md', '/escape.md', 'C:/escape.md', 'x\\escape.md',
                       '00-目录.md', 'build-report.json', 'images/x.md'):
            with self.subTest(unsafe=unsafe):
                chapter = copy.deepcopy(sections); chapter[0]['file'] = unsafe
                self.write_json(self.chapters_path, chapter)
                self.cli(book=True, ok=False)
                self.assertFalse(self.out.exists())
        self.write_json(self.chapters_path, sections)
        manifest['parts'][0]['id'] = '../escape'
        self.write_json(self.manifest_path, manifest)
        self.cli(book=True, ok=False)

    def test_figure_and_table_asset_escape_rejected_even_with_pdf(self):
        self.book(total=1, pdf=True)
        ocr = self.part('part01', [])
        for unsafe in ('../outside.png', '/outside.png', 'C:\\outside.png', '//host/image.png'):
            with self.subTest(unsafe=unsafe):
                with self.assertRaises(ValueError):
                    self.render([block('image', img_path=unsafe)], ocr_dir=ocr,
                                pdf_path=self.base / 'source.pdf')
                with self.assertRaises(ValueError):
                    self.render([block('table', table_body=f'<table><tr><td><img src="{unsafe}"></td></tr></table>')], ocr_dir=ocr)

    def test_only_preserves_other_chapters_index_report_review_and_assets(self):
        self.book()
        for i in (1, 2):
            self.part(f'part{i:02}', [block(text=f'正文{i}')])
        self.cli(book=True)
        other = (self.out / 'ch02.md').read_bytes()
        (self.out / 'ch02.md').write_bytes(other + b'\nmanual edit\n')
        (self.out / 'manual.md').write_text('人工文稿', encoding='utf-8')
        image = self.out / 'images' / 'ch02' / 'manual.png'
        image.parent.mkdir(parents=True)
        image.write_bytes(png('green'))
        self.part('part01', [block(text='更新第一章')])
        self.cli('--only', 'ch01', book=True)
        report = json.loads(self.text('build-report.json'))
        self.assertEqual([c['id'] for c in report['chapters']], ['ch01', 'ch02'])
        self.assertIn('ch02', report['review'])
        self.assertIn('ch02.md', self.text('00-目录.md'))
        self.assertEqual((self.out / 'ch02.md').read_bytes(), other + b'\nmanual edit\n')
        self.assertIn('更新第一章', self.text())
        self.assertTrue(image.exists())
        self.assertEqual(self.text('manual.md'), '人工文稿')

    def test_unowned_collision_never_overwritten_and_preflight_is_whole_batch(self):
        self.book()
        for i in (1, 2):
            self.part(f'part{i:02}', [block(text='正文')])
        self.out.mkdir()
        (self.out / 'ch02.md').write_text('人工文稿', encoding='utf-8')
        before = self.snapshot()
        for args in [(), ('--overwrite',)]:
            self.cli(*args, book=True, ok=False)
            self.assertEqual(self.snapshot(), before)

    def test_owned_manual_edits_require_overwrite(self):
        self.book(total=1)
        self.part('part01', [block(text='原正文')])
        self.cli(book=True)
        (self.out / 'ch01.md').write_text('人工修改', encoding='utf-8')
        before = self.snapshot()
        self.part('part01', [block(text='更新正文')])
        self.cli(book=True, ok=False)
        self.assertEqual(self.snapshot(), before)
        self.cli('--overwrite', book=True)
        self.assertIn('更新正文', self.text())
        self.assertIn('ch01.md', json.loads(self.text(OWNERSHIP))['files'])

    def test_clean_refused_and_unowned_index_report_preserved(self):
        self.part('ch01', [block(text='正文')])
        self.out.mkdir()
        for name in ['00-目录.md', 'build-report.json']:
            path = self.out / name
            path.write_text('人工内容', encoding='utf-8')
            before = self.snapshot()
            self.cli('--clean', ok=False)
            self.cli('--overwrite', ok=False)
            self.assertEqual(self.snapshot(), before)
            path.unlink()

    def test_regular_only_keeps_index_and_unknown_only_fails(self):
        for i in (1, 2):
            self.part(f'ch{i:02}', [block(text=f'正文{i}')])
        self.cli('--mode', 'book')
        self.part('ch01', [block(text='更新正文')])
        self.cli('--mode', 'book', '--only', 'ch01')
        self.assertIn('ch02.md', self.text('00-目录.md'))
        before = self.snapshot()
        self.cli('--mode', 'book', '--only', 'missing', ok=False)
        self.assertEqual(self.snapshot(), before)

    def test_nested_chapter_links_and_single_index(self):
        _, sections = self.book(total=1)
        sections[0]['file'] = 'chapters/ch01.md'
        self.write_json(self.chapters_path, sections)
        self.part('part01', [block('image', img_path='images/same.png')])
        self.cli(book=True)
        text = self.text('chapters/ch01.md')
        ref = self.refs(text)[0]
        self.assertTrue(ref.startswith('../images/'))
        self.assertTrue((self.out / 'chapters' / ref).is_file())
        self.assertIn('chapters/ch01.md', self.text('00-目录.md'))

    def test_output_parent_file_conflict_fails_without_changes(self):
        self.book(total=1)
        self.part('part01', [block('image', img_path='images/same.png')])
        self.out.mkdir()
        (self.out / 'images').write_text('人工文件', encoding='utf-8')
        before = self.snapshot()
        self.cli(book=True, ok=False)
        self.assertEqual(self.snapshot(), before)

    def test_pdf_total_mismatch_and_regular_ocr_bounds_fail(self):
        manifest, sections = self.book(total=2, pdf=True)
        manifest['pages'] = 1
        manifest['parts'] = manifest['parts'][:1]
        self.write_json(self.manifest_path, manifest)
        self.write_json(self.chapters_path, sections[:1])
        self.part('part01', [block(text='正文')])
        self.cli(book=True, ok=False)
        self.assertFalse(self.out.exists())
        self.part('part01', [block(text='越界', page=2)])
        self.cli('--pdf', str(self.base / 'source.pdf'), ok=False)
        self.assertFalse(self.out.exists())

    def test_in_page_geometric_order_sorting_restores_natural_flow(self):
        # Even if MinerU outputs the bottom figure first in the JSON array,
        # book mode sorts by vertical band so top prose appears first.
        ocr = self.part('part01', [])
        items = [
            block('image', bbox=[100, 600, 900, 900], img_path='images/same.png', image_caption=['图1 底部插图']),
            block(text='顶部先行正文', bbox=[100, 100, 900, 200]),
            block(text='底部后续正文', bbox=[100, 920, 900, 980]),
        ]
        text, _ = self.render(items, ocr_dir=ocr)
        top_pos = text.index('顶部先行正文')
        fig_pos = text.index('图1 底部插图')
        bot_pos = text.index('底部后续正文')
        self.assertLess(top_pos, fig_pos)
        self.assertLess(fig_pos, bot_pos)

    def test_single_number_subheading_distinguished_from_exercise(self):
        items = [
            block(text='1.1 正规小节', text_level=1),
            block(text='1. 链结构单元的化学组成', text_level=1),
            block(text='正文描述'),
            block(text='1. 计算这个数', text_level=1),
            block(text='复习题'),
            block(text='2. 晶格常数分析', text_level=1),
        ]
        text, _ = self.render(items)
        self.assertIn('## 1.1 正规小节', text)
        self.assertIn('### 1. 链结构单元的化学组成', text)
        self.assertIn('\n1. 计算这个数\n', text)
        self.assertNotIn('### 1. 计算这个数', text)
        # Inside exercise section, even noun headings remain list items
        self.assertIn('\n2. 晶格常数分析\n', text)
        self.assertNotIn('### 2. 晶格常数分析', text)

    def test_missing_nested_asset_and_fragment_fallback_fail(self):
        ocr = self.part('part01', [])
        with self.assertRaisesRegex(ValueError, '内嵌图片'):
            self.render([block('table', table_body='<table><tr><td><img src="missing.png"></td></tr></table>')], ocr_dir=ocr)
        with self.assertRaisesRegex(ValueError, '回退图'):
            self.render([block('image', img_path='images/same.png', bbox=[100, 200, 500, 700]),
                         block('image', img_path='missing.png', bbox=[300, 200, 800, 700])], ocr_dir=ocr)

    def test_process_repairs_inline_boundary_without_moving_prose(self):
        text, _ = self.render([block(text='正文（$x)$。')])
        self.assertIn('正文（$x$）。', text)

    def test_chapter_title_prefix_match_and_deduplication(self):
        items = [
            block(text='1 绪论', text_level=1),
            block(text='1.1 材料结构基础', text_level=2),
            block(text='正文描述'),
        ]
        text, _ = self.render(items, title='第1章 绪论')
        self.assertNotIn('## 1 绪论', text)
        self.assertIn('## 1.1 材料结构基础', text)

    def test_output_and_asset_symlinks_rejected_when_supported(self):
        outside = self.base / 'outside'
        outside.mkdir()
        try:
            self.out.symlink_to(outside, target_is_directory=True)
        except OSError:
            self.skipTest('Windows symlinks require permission')
        with self.assertRaises(ValueError):
            Publisher(self.out)
        (outside / 'image.png').write_bytes(png('red'))
        (self.raw / 'link.png').symlink_to(outside / 'image.png')
        with self.assertRaises(ValueError):
            asset_path(self.raw, 'link.png')

    def test_table_markdown_protects_inequalities(self):
        html = '<table><tr><td>Restoring force</td><td>$x << R_c$</td><td>$A < B$</td></tr></table>'
        res = build.table_markdown(html)
        self.assertIn('$x << R_c$', res)
        self.assertIn('$A < B$', res)

    def test_heading_preserves_math_formula(self):
        self.part('ch01', [
            block(text='1.3.2 $\\mathrm{CO}_2$ 辅助乙烷脱氢', text_level=2),
            block(text='正文内容')
        ])
        self.cli('--mode', 'book')
        text = self.text()
        self.assertIn('## 1.3.2 $\\mathrm{CO}_2$ 辅助乙烷脱氢', text)
        self.assertNotIn('CO_2', text)

    def test_ref_text_native_handling(self):
        self.part('ch01', [
            block(text='[1] Author. Title[J]. Journal, 2026.', kind='ref_text')
        ])
        self.cli('--mode', 'book')
        text = self.text()
        self.assertIn('[1] Author. Title[J]. Journal, 2026.', text)

    def test_prune_orphan_images(self):
        # Create an owned image that is referenced, and an owned image that is orphaned
        self.part('ch01', [
            block(kind='image', img_path='figure1.png', bbox=[100, 200, 300, 400],
                  image_caption=['图 1 测试']),
            block(text='正文')
        ])
        # Add a dummy figure1.png in ocr dir
        ch_dir = self.raw / 'ch01'
        (ch_dir / 'figure1.png').write_bytes(png('blue'))
        self.cli('--mode', 'book')
        # Verify figure cropped/stored
        state_dir = self.out.parent / f'.{self.out.name}.build'
        ownership_file = state_dir / OWNERSHIP
        data = json.loads(ownership_file.read_text(encoding='utf-8'))
        # Create an artificial orphan image in out/images/ch01/orphan.png and register it in ownership
        orphan_path = self.out / 'images' / 'ch01' / 'orphan.png'
        orphan_bytes = png('green')
        orphan_path.write_bytes(orphan_bytes)
        orphan_rel = 'images/ch01/orphan.png'
        data['files'][orphan_rel] = build.digest(orphan_bytes)
        ownership_file.write_text(json.dumps(data, indent=2), encoding='utf-8')
        self.assertTrue(orphan_path.exists())
        pruned = build._prune_orphan_images(self.out, state_dir)
        self.assertEqual(pruned, 1)
        self.assertFalse(orphan_path.exists())
        new_data = json.loads(ownership_file.read_text(encoding='utf-8'))
        self.assertNotIn(orphan_rel, new_data['files'])


    def test_auto_repairs_cli_applies_string_snapshot_notes_and_rebuilds(self):
        self.book(total=1, pdf=True)
        self.part('part01', [
            block('image', bbox=[100, 200, 400, 600], image_caption='Fig. S1. (a)', image_footnote='First note'),
            block('image', bbox=[450, 200, 800, 600], image_caption='(b)', image_footnote=['Second note']),
        ])
        self.cli('--auto-repairs', book=True)
        text = self.text()
        self.assertEqual(len(self.refs(text)), 1)
        self.assertIn('First note', text)
        self.assertIn('Second note', text)
        self.assertIn('## Figure S1', text)
        self.assertEqual(len(json.loads(self.text('build-report.json'))['repairs']), 2)
        before = self.snapshot()
        self.cli('--auto-repairs', book=True)
        self.assertEqual(self.snapshot(), before)

    def test_native_pdf_only_exact_standalone_region(self):
        store = build.FigureStore(self.base)
        with fitz.open() as doc:
            page = doc.new_page(width=200, height=100)
            page.insert_image(fitz.Rect(20, 20, 140, 80), stream=png('blue'), keep_proportion=False)
            self.assertIsNotNone(store.native_pdf_image(page, (.1, .2, .7, .8)))
            self.assertIsNone(store.native_pdf_image(page, (.1, .2, .8, .8)))
            self.assertIsNone(store.native_pdf_image(page, (.2, .3, .6, .7)))
            page.draw_line((20, 30), (130, 70))
            self.assertIsNone(store.native_pdf_image(page, (.1, .2, .7, .8)))

    def test_pptx_cropped_rotated_picture_uses_pdf(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        pic = slide.shapes.add_picture(io.BytesIO(png('green')), Inches(1), Inches(1), Inches(4), Inches(3))
        pic.crop_left = .2
        pic.rotation = 30
        source = self.base / 'source.pptx'
        prs.save(source)
        pages = build.load_source(source)
        self.assertEqual(pages[0]['pictures'], [])
        self.assertEqual(len(pages[0]['unsafe_pictures']), 1)
        self.book(total=1, pdf=True)
        _, stats = build.build_chapter([block('image')], pages, self.base / 'source.pdf',
                                       self.base / 'render/images/ch', 'Title', [], mode='book')
        self.assertEqual(stats['cropped_figures'], 1)
        self.assertEqual(stats.get('native_figures', 0), 0)
        with self.assertRaisesRegex(ValueError, 'requires a successful PDF crop'):
            build.build_chapter([block('image')], pages, None, self.base / 'unsafe/images/ch',
                                'Title', [], mode='book')

    def test_grouped_picture_is_not_extracted_as_raw_blob(self):
        from pptx import Presentation
        from pptx.util import Inches
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        group = slide.shapes.add_group_shape()
        group.shapes.add_picture(io.BytesIO(png('green')), Inches(1), Inches(1), Inches(4), Inches(3))
        source = self.base / 'grouped.pptx'
        prs.save(source)
        page = build.load_source(source)[0]
        self.assertEqual(page['pictures'], [])
        self.assertEqual(len(page['unsafe_pictures']), 1)

    def test_cross_page_tables_preserved_with_reported_range(self):
        review = []
        items = [block('table', page=i, table_caption='Table S1. Values',
                       table_body=f'<table><tr><td>row{i}</td></tr></table>') for i in (0, 1)]
        text, stats = build.build_chapter(items, None, None, self.base / 'tables/images/ch',
                                         'Title', review, mode='book')
        self.assertEqual(stats['tables'], 2)
        self.assertIn('row0', text)
        self.assertIn('row1', text)
        self.assertEqual(review[0]['page_range'], [1, 2])

    def test_prune_preserves_wiki_html_reference_and_encoded_assets(self):
        self.out.mkdir()
        image = self.out / 'images/ch1/a space.png'
        image.parent.mkdir(parents=True)
        image.write_bytes(png('red'))
        state = self.base / 'state'
        state.mkdir()
        ledger = state / OWNERSHIP
        for reference in ('![[a space.png]]', '<img src="images/ch1/a%20space.png">',
                          '![x][asset]\n\n[asset]: images/ch1/a%20space.png',
                          '[download](images/ch1/a%20space.png)', '![[missing.png]]'):
            (self.out / 'manual.md').write_text(reference, encoding='utf-8')
            ledger.write_text(json.dumps({'version': 1, 'files': {'images/ch1/a space.png': build.digest(image.read_bytes())}}), encoding='utf-8')
            with self.subTest(reference=reference):
                self.assertEqual(build._prune_orphan_images(self.out, state), 0)
                self.assertTrue(image.is_file())


if __name__ == '__main__':
    unittest.main()
